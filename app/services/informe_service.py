import re
import unicodedata
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Optional

from docx import Document

from app import rutas
from app.persistence.database import DB_PATH
from app.services import becario_service
from app.services.gestion_service import obtener_gestion_predeterminada

NOMBRE_PLANTILLA = "PLANTILLA.docx"
CARPETA_INFORMES = "informes"

CATEGORIAS = (
    ("excelencia_renov", "Excelencia Académica Renovación", 1),
    ("excelencia_nueva", "Excelencia Académica - Nueva", 2),
    ("economica_renov", "Beca Económica Social - Renovación", 3),
    ("convenio_renov", "Beca Convenio Interinstitucional - Renovación", 4),
    ("honorifica_renov", "Beca Honorífica Directorio - Renovación", 5),
    ("personal_admin", "Beca Personal Administrativo", 6),
    ("ministerio", "Beca Ministerio de Educación", 7),
    ("solicitudes_nuevas", "Beca Solicitudes Nuevas", 8),
)

INDICE_TABLA_ENCABEZADO = 0
INDICE_TABLA_RESUMEN = 9

COL_NUMERO = 0
COL_APELLIDOS = 1
COL_NOMBRES = 2
COL_CODIGO = 3
COL_CARRERA = 4
COL_PROMEDIO = 5
COL_TIPO = 6
COL_GESTION_INICIO = 7
COL_PORCENTAJE_ANT = 8
COL_PORCENTAJE_OTORGAR = 9

MESES_ES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
            "agosto", "septiembre", "octubre", "noviembre", "diciembre")


def _normalizar_tipo(tipo: str) -> str:
    base = unicodedata.normalize("NFKD", tipo or "")
    return "".join(c for c in base if not unicodedata.combining(c)).lower().strip()


def mapear_categoria(tipo_beca: str, estado: str,
                     carta_renovacion: bool) -> Optional[str]:
    if (estado or "").strip() == "Baja/Inactivo":
        return None
    if not carta_renovacion:
        return None
    tipo = _normalizar_tipo(tipo_beca)
    if "excelencia" in tipo:
        return "excelencia_renov"
    if "economica" in tipo and "social" in tipo:
        return "economica_renov"
    if "convenio" in tipo:
        return "convenio_renov"
    if "honor" in tipo:
        return "honorifica_renov"
    if "personal" in tipo or "administrativo" in tipo:
        return "personal_admin"
    if "ministerio" in tipo or "educacion" in tipo:
        return "ministerio"
    return None


def ruta_plantilla() -> Path:
    return rutas.assets_dir() / NOMBRE_PLANTILLA


def carpeta_informes() -> Path:
    return rutas.datos_dir() / CARPETA_INFORMES


def fecha_actual_es(fecha: Optional[datetime] = None) -> str:
    f = fecha or datetime.now()
    return f"{f.day} de {MESES_ES[f.month - 1]} de {f.year}"


def sugerir_nombre_acta(numero_acta: str, gestion: str) -> str:
    numero_limpio = re.sub(r'[\\/:*?"<>|\s]+', "", numero_acta or "").strip() or "SN"
    gestion_limpia = re.sub(r'[\\/:*?"<>|\s]+', "", gestion or "").strip()
    return f"Acta_N{numero_limpio}_{gestion_limpia}.docx"


def _nombre_unico(carpeta: Path, nombre: str) -> Path:
    base = re.sub(r'[\\/:*?"<>|]+', "", nombre).strip() or "Acta.docx"
    if not base.lower().endswith(".docx"):
        base += ".docx"
    destino = carpeta / base
    contador = 2
    while destino.exists():
        destino = carpeta / f"{base[:-5]}_v{contador}.docx"
        contador += 1
    return destino


def _escribir_celda(celda, texto: str):
    texto = texto or ""
    parrafos = celda.paragraphs
    if not parrafos:
        celda.text = texto
        return
    primero = parrafos[0]
    if primero.runs:
        primero.runs[0].text = texto
        for run in primero.runs[1:]:
            run.text = ""
    else:
        primero.text = texto
    for extra in parrafos[1:]:
        for run in extra.runs:
            run.text = ""


def _reemplazar_total_narrativo(doc, total: int) -> int:
    patron = re.compile(r"(un total de\s+)\d+(\s+becas)")
    reemplazados = 0

    def _procesar(parrafo) -> bool:
        completo = "".join((run.text or "") for run in parrafo.runs)
        nuevo, cantidad = patron.subn(lambda m: f"{m.group(1)}{total}{m.group(2)}",
                                      completo)
        if not cantidad:
            return False
        for run in parrafo.runs:
            if "66" in (run.text or ""):
                run.text = (run.text or "").replace("66", str(total))
        if "".join((run.text or "") for run in parrafo.runs) != nuevo:
            parrafo.runs[0].text = nuevo
            for run in parrafo.runs[1:]:
                run.text = ""
        return True

    for parrafo in doc.paragraphs:
        if parrafo.runs and _procesar(parrafo):
            reemplazados += 1
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                for parrafo in celda.paragraphs:
                    if parrafo.runs and _procesar(parrafo):
                        reemplazados += 1
    return reemplazados


def _rellenar_tabla_categoria(tabla, filas: list[tuple]) -> int:
    faltantes = 0
    elemento = tabla._tbl
    plantilla = deepcopy(tabla.rows[1]._tr)
    for fila in list(tabla.rows[1:]):
        elemento.remove(fila._tr)
    for numero, dato in enumerate(filas, start=1):
        apellidos, nombres, codigo, carrera, tipo, inicio, porc_ant = dato
        if not inicio:
            faltantes += 1
        if not porc_ant:
            faltantes += 1
        elemento.append(deepcopy(plantilla))
        celdas = tabla.rows[-1].cells
        _escribir_celda(celdas[COL_NUMERO], str(numero))
        _escribir_celda(celdas[COL_APELLIDOS], apellidos)
        _escribir_celda(celdas[COL_NOMBRES], nombres)
        _escribir_celda(celdas[COL_CODIGO], codigo)
        _escribir_celda(celdas[COL_CARRERA], carrera)
        _escribir_celda(celdas[COL_PROMEDIO], "")
        _escribir_celda(celdas[COL_TIPO], tipo)
        _escribir_celda(celdas[COL_GESTION_INICIO], inicio)
        _escribir_celda(celdas[COL_PORCENTAJE_ANT], porc_ant)
        _escribir_celda(celdas[COL_PORCENTAJE_OTORGAR], "")
    return faltantes


def generar_acta(numero_acta: str, nombre_archivo: Optional[str] = None,
                 db_path: Path = DB_PATH,
                 fecha: Optional[datetime] = None) -> dict:
    numero = (numero_acta or "").strip()
    if not numero:
        raise ValueError("El número de acta es obligatorio.")
    gestion = obtener_gestion_predeterminada(db_path)
    plantilla = ruta_plantilla()
    if not plantilla.is_file():
        raise FileNotFoundError(f"No se encontró la plantilla: {plantilla}")

    grupos: dict[str, list] = {clave: [] for clave, _, _ in CATEGORIAS}
    omitidos: list[tuple[str, str, str]] = []
    for becario, seg in becario_service.listar_para_panel(db_path):
        carta = bool(seg is not None and seg.carta_renovacion)
        clave = mapear_categoria(becario.tipo_beca, becario.estado, carta)
        if clave is None:
            if (becario.estado or "").strip() == "Baja/Inactivo":
                motivo = "Baja/Inactivo"
            elif not carta:
                motivo = "sin carta de renovación"
            else:
                motivo = f"tipo sin tabla de renovación ({becario.tipo_beca or '—'})"
            omitidos.append((becario.apellidos or "", becario.nombres or "", motivo))
            continue
        grupos[clave].append((
            becario.apellidos or "", becario.nombres or "",
            becario.codigo_estudiante or "", becario.carrera or "",
            becario.tipo_beca or "", becario.gestion_ingreso or "",
            (seg.porcentaje_anterior or "") if seg is not None else "",
        ))

    doc = Document(str(plantilla))
    encabezado = doc.tables[INDICE_TABLA_ENCABEZADO].rows[0].cells
    _escribir_celda(encabezado[0], f"Fecha: {fecha_actual_es(fecha)}")
    _escribir_celda(encabezado[1], f"Acta N°: {numero}")
    _escribir_celda(encabezado[2], f"Gestión: {gestion}")

    faltantes = 0
    conteos = []
    for clave, etiqueta, indice in CATEGORIAS:
        if grupos[clave]:
            faltantes += _rellenar_tabla_categoria(doc.tables[indice], grupos[clave])
        conteos.append((clave, etiqueta, len(grupos[clave])))

    total = sum(cantidad for _, _, cantidad in conteos)
    resumen = doc.tables[INDICE_TABLA_RESUMEN]
    for i, (_clave, _etiqueta, cantidad) in enumerate(conteos, start=1):
        _escribir_celda(resumen.rows[i].cells[2], str(cantidad))
    _escribir_celda(resumen.rows[len(conteos) + 1].cells[2], str(total))
    _reemplazar_total_narrativo(doc, total)

    celdas_vacias = total * 2 + faltantes
    destino = carpeta_informes()
    destino.mkdir(parents=True, exist_ok=True)
    ruta = _nombre_unico(destino, nombre_archivo or sugerir_nombre_acta(numero, gestion))
    doc.save(str(ruta))
    return {"ruta": str(ruta), "nombre_archivo": ruta.name, "numero_acta": numero,
            "gestion": gestion, "conteos": conteos, "total": total,
            "celdas_vacias": celdas_vacias, "omitidos": omitidos}


def listar_informes() -> list[dict]:
    carpeta = carpeta_informes()
    if not carpeta.is_dir():
        return []
    archivos = []
    for ruta in carpeta.glob("*.docx"):
        try:
            stat = ruta.stat()
        except OSError:
            continue
        archivos.append({"nombre": ruta.name, "ruta": str(ruta),
                         "tamano_bytes": stat.st_size,
                         "modificado": datetime.fromtimestamp(stat.st_mtime)})
    archivos.sort(key=lambda a: a["modificado"], reverse=True)
    return archivos


def eliminar_informe(nombre: str) -> bool:
    seguro = Path(nombre or "").name
    if not seguro or seguro in (".", ".."):
        raise ValueError("Nombre de informe no válido.")
    ruta = carpeta_informes() / seguro
    if not ruta.is_file():
        return False
    ruta.unlink()
    return True
